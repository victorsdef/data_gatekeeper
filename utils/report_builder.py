"""
utils/report_builder.py
Genera el reporte Excel de errores de validación.
Hoja 1 (Resumen): metadata de la carga + conteo por tipo de error + leyenda.
Hoja 2 (Detalle):  tabla completa con filas coloreadas por categoría de error.
"""
from __future__ import annotations

import io
from collections import Counter
from datetime import datetime
from typing import Any, Dict

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from dg_validators.engine import ValidationResult


# ------------------------------------------------------------------
# Paleta de colores por categoría de error
# ------------------------------------------------------------------
_RULE_COLOR: Dict[str, str] = {
    "columna_requerida":   "FFC7CE",  # rojo claro   → estructural
    "columna_no_esperada": "FFC7CE",
    "archivo_vacio":       "FFC7CE",
    "tipo_de_dato":        "FFEB9C",  # amarillo      → tipo / nulabilidad
    "no_nulo":             "FFEB9C",
    "dominio":             "FCE4D6",  # naranja claro → dominio / rango
    "valor_mínimo":        "FCE4D6",
    "valor_máximo":        "FCE4D6",
    "longitud_mínima":     "FCE4D6",
    "longitud":            "FCE4D6",
    "error_inesperado":    "FFC7CE",
}

_RULE_LABEL: Dict[str, str] = {
    "columna_requerida":   "Columna faltante",
    "columna_no_esperada": "Columna extra no esperada",
    "archivo_vacio":       "Archivo vacío",
    "tipo_de_dato":        "Tipo de dato incorrecto",
    "no_nulo":             "Valor nulo no permitido",
    "dominio":             "Valor fuera de dominio",
    "valor_mínimo":        "Valor bajo el mínimo permitido",
    "valor_máximo":        "Valor sobre el máximo permitido",
    "longitud_mínima":     "Longitud insuficiente",
    "longitud":            "Longitud fuera de rango permitido",
    "error_inesperado":    "Error inesperado",
}

_LEGEND = [
    ("FFC7CE", "Error estructural — columnas faltantes / extras / archivo vacío"),
    ("FFEB9C", "Error de tipo de dato o valor nulo no permitido"),
    ("FCE4D6", "Error de dominio, rango o longitud"),
]

# Colores de marca
_BRAND   = "534AB7"
_SUCCESS = "1D9E75"
_ERROR   = "C00000"
_GRAY    = "F4F5F9"
_DARK    = "374151"


# ------------------------------------------------------------------
# Helpers de estilo
# ------------------------------------------------------------------
def _fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color)

def _font(bold: bool = False, color: str = "000000", size: int = 10) -> Font:
    return Font(bold=bold, color=color, name="Calibri", size=size)

def _center() -> Alignment:
    return Alignment(horizontal="center", vertical="center", wrap_text=True)

def _left() -> Alignment:
    return Alignment(horizontal="left", vertical="center", wrap_text=True)


# ------------------------------------------------------------------
# Función pública
# ------------------------------------------------------------------
def build_error_report(
    result: ValidationResult,
    catalog: Dict[str, Any],
    filename: str,
    username: str,
) -> bytes:
    """Retorna los bytes de un archivo .xlsx con el reporte de validación."""
    wb = openpyxl.Workbook()
    _build_summary_sheet(wb, result, catalog, filename, username)
    _build_detail_sheet(wb, result)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------------
# Hoja 1: Resumen
# ------------------------------------------------------------------
def _build_summary_sheet(
    wb: openpyxl.Workbook,
    result: ValidationResult,
    catalog: Dict[str, Any],
    filename: str,
    username: str,
) -> None:
    ws = wb.active
    ws.title = "Resumen"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 42

    row = 1

    # Título principal
    ws.merge_cells(f"A{row}:B{row}")
    c = ws[f"A{row}"]
    c.value = "Data Gatekeeper — Reporte de Validación"
    c.font = _font(bold=True, color="FFFFFF", size=14)
    c.fill = _fill(_BRAND)
    c.alignment = _center()
    ws.row_dimensions[row].height = 32
    row += 1

    # Metadata
    meta = [
        ("Archivo cargado",    filename),
        ("Catálogo",           catalog.get("nombre", "—")),
        ("Tabla destino",      catalog.get("tabla_destino", "—")),
        ("Estrategia",         catalog.get("estrategia", "—").upper()),
        ("Destino",            catalog.get("destino", "—").upper()),
        ("Usuario",            username),
        ("Fecha y hora",       datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("Filas revisadas",    f"{result.rows_checked:,}"),
        ("Columnas revisadas", str(result.cols_checked)),
    ]
    for label, value in meta:
        lc = ws.cell(row=row, column=1, value=label)
        lc.font = _font(bold=True, size=10)
        lc.fill = _fill(_GRAY)
        lc.alignment = _left()

        vc = ws.cell(row=row, column=2, value=value)
        vc.font = _font(size=10)
        vc.alignment = _left()

        ws.row_dimensions[row].height = 18
        row += 1

    row += 1  # espacio

    # Banner resultado
    ws.merge_cells(f"A{row}:B{row}")
    c = ws[f"A{row}"]
    if result.success:
        c.value = "✓  VALIDACIÓN EXITOSA — Sin errores encontrados"
        c.fill = _fill(_SUCCESS)
    else:
        c.value = f"✗  VALIDACIÓN FALLIDA — {result.error_count} error(es) encontrado(s)"
        c.fill = _fill(_ERROR)
    c.font = _font(bold=True, color="FFFFFF", size=12)
    c.alignment = _center()
    ws.row_dimensions[row].height = 28
    row += 1

    if not result.errors:
        return

    row += 1  # espacio

    # Encabezado de resumen por tipo
    ws.merge_cells(f"A{row}:B{row}")
    h = ws[f"A{row}"]
    h.value = "Errores por tipo"
    h.font = _font(bold=True, color="FFFFFF", size=11)
    h.fill = _fill(_DARK)
    h.alignment = _center()
    ws.row_dimensions[row].height = 22
    row += 1

    counts = Counter(e.regla for e in result.errors)
    for regla, count in sorted(counts.items(), key=lambda x: -x[1]):
        color = _RULE_COLOR.get(regla, "FFFFFF")
        lc = ws.cell(row=row, column=1, value=_RULE_LABEL.get(regla, regla))
        lc.fill = _fill(color)
        lc.font = _font(size=10)
        lc.alignment = _left()

        vc = ws.cell(row=row, column=2, value=count)
        vc.fill = _fill(color)
        vc.font = _font(bold=True, size=10)
        vc.alignment = _center()

        ws.row_dimensions[row].height = 16
        row += 1

    row += 1  # espacio

    # Leyenda
    ws.merge_cells(f"A{row}:B{row}")
    h = ws[f"A{row}"]
    h.value = "Leyenda de colores (hoja Detalle)"
    h.font = _font(bold=True, size=10)
    h.alignment = _left()
    ws.row_dimensions[row].height = 16
    row += 1

    for hex_c, desc in _LEGEND:
        swatch = ws.cell(row=row, column=1, value="")
        swatch.fill = _fill(hex_c)

        label = ws.cell(row=row, column=2, value=desc)
        label.font = _font(size=10)
        label.alignment = _left()

        ws.row_dimensions[row].height = 14
        row += 1


# ------------------------------------------------------------------
# Hoja 2: Detalle de errores
# ------------------------------------------------------------------
def _build_detail_sheet(wb: openpyxl.Workbook, result: ValidationResult) -> None:
    ws = wb.create_sheet("Detalle de errores")
    ws.sheet_view.showGridLines = False

    headers    = ["Fila en archivo", "Columna", "Valor encontrado", "Tipo de error", "Descripción del problema"]
    col_widths = [16,                22,         30,                 26,              60]

    for i, w in enumerate(col_widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # Fila de encabezados
    for col, h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=col, value=h)
        c.font = _font(bold=True, color="FFFFFF", size=10)
        c.fill = _fill(_BRAND)
        c.alignment = _center()
    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "A2"  # congela encabezados al hacer scroll

    if not result.errors:
        ws.cell(row=2, column=1, value="Sin errores").font = _font(bold=True, color=_SUCCESS)
        return

    for row_i, err in enumerate(result.errors, start=2):
        color = _RULE_COLOR.get(err.regla, "FFFFFF")
        values = [
            err.fila if err.fila > 0 else "—",
            err.columna,
            str(err.valor),
            _RULE_LABEL.get(err.regla, err.regla),
            err.detalle,
        ]
        for col, val in enumerate(values, start=1):
            c = ws.cell(row=row_i, column=col, value=val)
            c.fill = _fill(color)
            c.font = _font(size=10)
            c.alignment = _left()
        ws.row_dimensions[row_i].height = 15
